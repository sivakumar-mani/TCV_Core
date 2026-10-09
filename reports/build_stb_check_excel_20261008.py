import json
import sys
import re
from pathlib import Path
from collections import Counter
from datetime import datetime, timezone, timedelta
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo

root = Path(__file__).resolve().parent
report_tag = sys.argv[1] if len(sys.argv) > 1 else 'stb_check_20261008'
if not re.fullmatch(r'stb_check_\d{8}(?:_batch\d+)?', report_tag):
    raise ValueError('Invalid report name')
data = json.loads((root / (report_tag + '_results.json')).read_text())
wb = Workbook()
summary = wb.active
summary.title = 'STB Summary'
summary.append(['S.No', 'Requested STB Number', 'Result', 'Active Customer', 'Active Approved STB Assignment', 'Network', 'Customer IDs', 'Customer Names', 'Customer Status', 'STB Status', 'History Matches', 'Active Assignment Count', 'Notes'])
counts = Counter()
for index, number in enumerate(data['input'], 1):
    matched = [r for r in data['rows'] if r['stb_no'].strip().upper() == number.upper()]
    latest = {}
    for row in matched:
        if row['stb_approval'] == 'APPROVED':
            key = row['cable_customer_id']
            if key not in latest or int(row['customer_stb_id']) > int(latest[key]['customer_stb_id']):
                latest[key] = row
    selected = list(latest.values()) or matched
    customers = {r['cable_customer_id']: r for r in selected}
    active_customers = [r for r in customers.values() if r['customer_status'] == 'ACTIVE' and r['customer_approval'] == 'APPROVED']
    active_assignments = [r for r in latest.values() if r['customer_status'] == 'ACTIVE' and r['customer_approval'] == 'APPROVED' and r['stb_status'] == 'ACTIVE']
    master = [r for r in data['masters'] if r['stb_number'].strip().upper() == number.upper()]
    notes = []
    if 'E+' in number.upper():
        result = 'AMBIGUOUS INPUT'
        notes.append('Scientific notation may have lost precision/leading zeros. Supply the original STB number as text; no guessed lookup performed.')
    elif active_assignments:
        result = 'ACTIVE MATCH'
    elif matched:
        result = 'MATCHED - NOT ACTIVE ASSIGNMENT'
    elif master:
        result = 'MASTER ONLY - NO CUSTOMER'
    else:
        result = 'NOT FOUND'
    if len(active_assignments) > 1:
        notes.append('Multiple active customer associations found; inspect Match Details.')
    if number.upper().startswith('000083') and len(number) != 16:
        notes.append('Length differs from the 16-character STB numbers in this list; checked exactly as supplied. Confirm original number if not found.')
    if matched and not latest:
        notes.append('No approved STB history for this number; showing pending/rejected rows.')
    if master:
        notes.append('Also found in STB master: ' + ', '.join(str(r['stb_master_id']) for r in master))
    counts[result] += 1
    def joined(field):
        return '; '.join(sorted({str(r.get(field) or '') for r in selected} - {''}))
    networks = '; '.join(sorted({r.get('network_name') or r.get('network_code') or r.get('network_type') or 'Unknown' for r in selected}))
    summary.append([index, number, result, 'YES' if active_customers else 'NO' if matched else 'UNKNOWN', 'YES' if active_assignments else 'NO' if matched else 'UNKNOWN', networks, joined('cable_customer_id'), joined('full_name'), joined('customer_status'), joined('stb_status'), len(matched), len(active_assignments), ' '.join(notes)])

details = wb.create_sheet('Match Details')
fields = ['stb_no', 'customer_stb_id', 'stb_status', 'stb_approval', 'stb_type', 'installed_date', 'cable_customer_id', 'customer_code', 'legacy_customer_no', 'network_customer_no', 'full_name', 'customer_status', 'customer_approval', 'network_code', 'network_name', 'network_type', 'latest_approved_stb_id']
details.append(['Requested STB Number'] + [f.replace('_', ' ').title() for f in fields] + ['Latest Approved STB Row for Customer'])
for number in data['input']:
    for row in data['rows']:
        if row['stb_no'].strip().upper() == number.upper():
            details.append([number] + [row.get(f) for f in fields] + ['YES' if str(row['customer_stb_id']) == str(row['latest_approved_stb_id']) else 'NO'])
masters = wb.create_sheet('STB Master Matches')
masters.append(['STB Master ID', 'STB Number'])
for row in data['masters']:
    masters.append([row['stb_master_id'], row['stb_number']])
info = wb.create_sheet('Report Notes')
info.append(['Item', 'Value'])
info.append(['Checked at (India)', datetime.fromisoformat(data['checked_at'].replace('Z', '+00:00')).astimezone(timezone(timedelta(hours=5, minutes=30))).strftime('%d-%m-%Y %H:%M:%S IST')])
info.append(['Source', f"{data['connection_type']}: {data['database']} on port {data['port']}. Production has not been verified."])
info.append(['Inputs', len(data['input'])])
info.append(['Lookup', 'Case-insensitive exact STB-number matching after trimming surrounding spaces. Leading zeros retained.'])
info.append(['Active customer', 'Customer status ACTIVE and customer approval APPROVED.'])
info.append(['Active approved assignment', 'For each matched STB number/customer pair, select its latest APPROVED STB row; both STB and customer must be ACTIVE and customer APPROVED.'])
info.append(['History', 'Match Details contains all matching history rows, including inactive, pending and rejected. Latest Approved STB Row for Customer indicates the overall customer latest row; this is informational and does not prove other STBs are inactive.'])
info.append(['Network', 'Customer network from cable_network_master; fallback to customer.network_type. STB vendor/MSO is not the customer network.'])
info.append(['Ambiguous numbers', ', '.join(n for n in data['input'] if 'E+' in n.upper()) + ': scientific notation cannot identify an exact STB number reliably. Flagged rather than expanded or guessed.'])
info.append(['Changes', 'SELECT-only lookup inside a read-only transaction. No database records changed.'])
for name, count in sorted(counts.items()):
    info.append([name, count])

for ws in wb:
    ws.freeze_panes = 'C2' if ws == summary else 'A2'
    for cell in ws[1]:
        cell.fill = PatternFill('solid', fgColor='17365D')
        cell.font = Font(color='FFFFFF', bold=True)
        cell.alignment = Alignment(wrap_text=True)
    ws.row_dimensions[1].height = 32
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = 's'
                cell.number_format = '@'
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    for col in ws.columns:
        max_len = max(len(str(c.value or '')) for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(max(max_len + 2, 14), 48)
    if ws.max_row > 1:
        tab = Table(displayName=ws.title.replace(' ', ''), ref=ws.dimensions)
        tab.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
        ws.add_table(tab)
for row in summary.iter_rows(min_row=2):
    color = 'E2F0D9' if row[2].value == 'ACTIVE MATCH' else 'FFF2CC' if row[2].value == 'AMBIGUOUS INPUT' else 'FCE4D6'
    row[2].fill = PatternFill('solid', fgColor=color)
info.column_dimensions['B'].width = 110
filename = root / ('STB_Customer_Active_Network_Check_' + report_tag.removeprefix('stb_check_') + '.xlsx')
wb.save(filename)
check = load_workbook(filename)
assert check['STB Summary'].max_row == len(data['input']) + 1
assert [check['STB Summary'].cell(i + 2, 2).value for i in range(len(data['input']))] == data['input']
assert check['Match Details'].max_row == len(data['rows']) + 1
assert sum(counts.values()) == len(data['input'])
print(json.dumps({'file': str(filename), 'inputs': len(data['input']), 'results': dict(counts), 'history_rows': len(data['rows'])}, indent=2))
